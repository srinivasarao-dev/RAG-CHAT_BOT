"""Chat with your files: text extraction, chunking and local embeddings for retrieval (RAG).

Embeddings run locally with fastembed (BAAI/bge-small-en-v1.5, 384 dims, ~70 MB, no API key).
Vectors are stored next to the chunks (Supabase pgvector or SQLite) and searched by cosine similarity.
"""
from __future__ import annotations

import io
import os
import re
import threading
from pathlib import Path

import numpy as np

EMBED_MODEL = "BAAI/bge-small-en-v1.5"
EMBED_DIM = 384
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 2_000_000
CHUNK_CHARS = 900
CHUNK_OVERLAP = 150

IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp", ".gif": "image/gif"}
TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".log",
    ".xml", ".sql", ".py", ".ipynb", ".js", ".jsx", ".ts", ".tsx", ".java", ".kt", ".go", ".rs", ".rb", ".php", ".c",
    ".h", ".cpp", ".hpp", ".cs", ".swift", ".scala", ".sh", ".bat", ".ps1", ".css", ".scss", ".vue", ".svelte", ".r",
    ".dart", ".lua", ".env.example", ".dockerfile", ".gradle", ".tex", ".rst",
}

_model = None
_model_lock = threading.Lock()


def _embedder():
    global _model
    with _model_lock:
        if _model is None:
            from fastembed import TextEmbedding

            # FASTEMBED_CACHE_PATH lets a Docker image ship the model instead of downloading it on start
            _model = TextEmbedding(EMBED_MODEL, cache_dir=os.getenv("FASTEMBED_CACHE_PATH") or None)
        return _model


def warm_up() -> None:
    """Load the embedding model in the background so the first upload isn't slow."""
    threading.Thread(target=lambda: _embedder(), daemon=True).start()


def embed_texts(texts: list[str]) -> np.ndarray:
    vecs = np.array(list(_embedder().embed(texts, batch_size=64)), dtype=np.float32)
    return vecs / np.maximum(np.linalg.norm(vecs, axis=1, keepdims=True), 1e-9)


def embed_query(text: str) -> np.ndarray:
    v = np.array(list(_embedder().query_embed([text]))[0], dtype=np.float32)
    return v / max(float(np.linalg.norm(v)), 1e-9)


class UnsupportedFile(ValueError):
    pass


def kind_of(filename: str) -> str:
    ext = Path(filename.lower()).suffix
    if ext in IMAGE_TYPES:
        return "image"
    return "document"


def extract_text(filename: str, data: bytes) -> str:
    """Plain text from an uploaded file. Raises UnsupportedFile for formats we can't read."""
    name = filename.lower()
    ext = Path(name).suffix
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = []
        for i, page in enumerate(reader.pages, 1):
            text = (page.extract_text() or "").strip()
            if text:
                pages.append(f"[Page {i}]\n{text}")
        if not pages:
            raise UnsupportedFile("This PDF has no extractable text (it may be a scanned image).")
        return "\n\n".join(pages)
    if ext == ".docx":
        import docx

        doc = docx.Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    if ext in (".html", ".htm"):
        try:
            import trafilatura

            text = trafilatura.extract(data.decode("utf-8", "replace"), include_tables=True)
            if text:
                return text
        except Exception:
            pass
        from bs4 import BeautifulSoup

        return BeautifulSoup(data, "html.parser").get_text("\n")
    if ext == ".ipynb":
        import json

        nb = json.loads(data.decode("utf-8", "replace"))
        cells = []
        for cell in nb.get("cells", []):
            src = "".join(cell.get("source", []))
            cells.append(f"# [{cell.get('cell_type', 'cell')}]\n{src}")
        return "\n\n".join(cells)
    if ext in TEXT_EXTS or name.endswith((".env.example", "dockerfile", "makefile", "readme")) or not ext:
        for enc in ("utf-8", "utf-16", "latin-1"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
    raise UnsupportedFile(f"Can't read '{ext or filename}' files yet. Try PDF, Word (.docx), text, Markdown, CSV, HTML or code.")


def chunk_text(text: str, size: int = CHUNK_CHARS, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Split on paragraph boundaries into ~size-char chunks with some overlap for context."""
    text = re.sub(r"\n{3,}", "\n\n", text.replace("\r\n", "\n")).strip()
    if not text:
        return []
    paras, chunks, cur = text.split("\n\n"), [], ""
    for p in paras:
        while len(p) > size:  # very long paragraph: hard-split it
            head, p = p[:size], p[size - overlap:]
            if cur:
                chunks.append(cur)
                cur = ""
            chunks.append(head)
        if len(cur) + len(p) + 2 <= size:
            cur = f"{cur}\n\n{p}" if cur else p
        else:
            if cur:
                chunks.append(cur)
            tail = cur[-overlap:] if cur else ""
            cur = f"{tail}\n\n{p}".strip() if tail else p
    if cur:
        chunks.append(cur)
    return chunks


def process_file(filename: str, data: bytes) -> dict:
    """Extract, chunk and embed a document. Returns {text, chunks, embeddings}."""
    if len(data) > MAX_FILE_BYTES:
        raise UnsupportedFile(f"File is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB.")
    text = extract_text(filename, data)[:MAX_TEXT_CHARS].strip()
    if not text:
        raise UnsupportedFile("No text found in this file.")
    chunks = chunk_text(text)
    return {"text": text, "chunks": chunks, "embeddings": embed_texts(chunks)}


def top_k(query_vec: np.ndarray, rows: list[dict], k: int = 6) -> list[dict]:
    """Rank chunk rows ({'embedding': np.ndarray, ...}) by cosine similarity."""
    if not rows:
        return []
    mat = np.stack([r["embedding"] for r in rows])
    scores = mat @ query_vec
    order = np.argsort(-scores)[:k]
    return [{**{k2: v for k2, v in rows[i].items() if k2 != "embedding"}, "score": float(scores[i])} for i in order]
